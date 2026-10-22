import os
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
p=Path(__file__).resolve().parent / 'qa-artifacts' / 'chart-real.xlsx'
p.parent.mkdir(parents=True,exist_ok=True)
w=Workbook();s=w.active;s.title='売上'
s.column_dimensions['A'].width=22
s['A1']='商品';s['B1']='数量';s['A2']='鉛筆';s['B2']=3;s['C2']='=B2*2'
s['A3']='ノート';s['B3']=7
s.freeze_panes='A2'
s['A1'].fill=PatternFill('solid',fgColor='1D67AA')
s['A1'].font=Font(bold=True,color='FFFFFF')
v=DataValidation(type='list',formula1='"A,B,C"');s.add_data_validation(v);v.add(s['B4'])
s.conditional_formatting.add('B2:B3',CellIsRule(operator='greaterThan',formula=['5'],fill=PatternFill('solid',fgColor='FF3300')))
c=BarChart();c.type='col';c.title='数量比較';c.add_data(Reference(s,min_col=2,min_row=1,max_row=3),titles_from_data=True);c.set_categories(Reference(s,min_col=1,min_row=2,max_row=3));s.add_chart(c,'E2')
other=w.create_sheet('備考');other['A1']='グラフを維持'
w.save(p)
print(p,p.stat().st_size)
